"""skills/media.py — video transcription and document text extraction."""

import subprocess
import tempfile
from pathlib import Path
from langchain_core.tools import tool


def _extract_audio(video_path: str, audio_out: str = "") -> str:
    """
    Extract audio from a video file using ffmpeg.

    Args:
        video_path: Absolute path to the input video file.
        audio_out: Output path for the extracted WAV file. Auto-set if empty.

    Returns:
        str: Path to the extracted audio file.

    Raises:
        RuntimeError: If ffmpeg fails.
    """
    if not audio_out:
        audio_out = str(Path(tempfile.gettempdir()) / "extracted_audio.wav")
    result = subprocess.run(
        ["ffmpeg", "-y", "-i", video_path, "-vn", "-ar", "16000", "-ac", "1", audio_out],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        raise RuntimeError(f"ffmpeg error: {result.stderr}")
    return audio_out


def _extract_text(file_path: str) -> str:
    """
    Extract plain text from a file based on its extension.

    Args:
        file_path: Absolute path to the file.

    Returns:
        str: Extracted text content.
    """
    path = Path(file_path).resolve()
    suffix = path.suffix.lower()

    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix in {".xlsx", ".xls"}:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join(str(c) for c in row if c is not None))
        return "\n".join(lines)

    return path.read_text(encoding="utf-8", errors="replace")


@tool
def transcribe_video(video_path: str) -> str:
    """
    Transcribe a local video or audio file to text using faster-whisper.
    Fully local — no cloud API required.

    Args:
        video_path: Absolute path to the video/audio file.

    Returns:
        str: Full transcribed text, or error message.
    """
    try:
        from faster_whisper import WhisperModel
        audio_path = _extract_audio(video_path)
        model = WhisperModel("base", device="cpu", compute_type="int8")
        segments, _ = model.transcribe(audio_path)
        return " ".join(seg.text for seg in segments)
    except ImportError:
        return "faster-whisper not installed. Run: pip install faster-whisper"
    except Exception as e:
        return f"Transcription error: {e}"


@tool
def extract_document_text(file_path: str) -> str:
    """
    Extract and return the full text content of a document for summarization.
    Supports .txt, .md, .py, .docx, .pdf, .xlsx files.

    The agent LLM will summarize the returned text — this tool only extracts it.

    Args:
        file_path: Absolute path to the document file.

    Returns:
        str: Full extracted text content (up to 6000 chars).
    """
    path = Path(file_path).resolve()
    if not path.exists():
        return f"Error: File not found: {file_path}"
    try:
        content = _extract_text(file_path)
        return content[:6000]
    except Exception as e:
        return f"Error extracting text: {e}"
